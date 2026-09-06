"""Vergleicht Sprachmodelle und Kontextlängen an einem festen Testfall.

Für jede Kombination aus Modell und Kontextlänge läuft immer derselbe Ablauf,
genau wie ihn eine ratsuchende Person durchliefe:

    1. die fünf Unterlagen des Beispielfalls einlesen
    2. daraus den Suchindex aufbauen
    3. die Aufgabe "Differenzanalyse" stellen
    4. die Antwort gegen die bekannte Wahrheit bewerten

Der Beispielfall ist eigens dafür gebaut: Hannelore Brandt hat 42,5 gewichtete
Punkte und damit Pflegegrad 2; bis Pflegegrad 3 fehlen genau 5 Punkte, und die
Belege dafür liegen verteilt in Pflegetagebuch, Befundbericht und
Medikationsplan. Damit ist prüfbar, ob ein Modell sie findet - statt bloß, ob
die Antwort gut klingt.

WICHTIG - die Grenzen im Code wachsen mit:
Bliebe MAX_CONTEXT_CHARS bei 18.000 Zeichen, wäre der Prompt bei 15.000 und bei
50.000 Token identisch. Man würde ausschließlich Wartezeit messen und daraus
schließen, mehr Kontext bringe nichts. Deshalb rechnet ``grenzen_setzen`` das
Budget je Spalte neu aus.

Aufruf:
    python vergleich_modelle.py                  – alle Kombinationen
    python vergleich_modelle.py --nur-schaetzen  – nur prüfen, was in den Speicher passt
"""

from __future__ import annotations

import argparse
import io
import json
import os
import subprocess
import sys
import time
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import List

sys.stdout = io.TextIOWrapper(
    sys.stdout.buffer, encoding="utf-8", errors="replace", line_buffering=True
)

LMS = Path.home() / ".lmstudio" / "bin" / "lms.exe"
UNTERLAGEN = Path(
    r"C:\Users\Timo\Wirtschaftsinformatik Hildesheim\Semester 2"
    r"\IT-Studienprojekt\beispielfall"
)

# Vier Modelle in einer Gewichtsklasse (6 bis 9 GB in Q4_K_M), damit sie den
# Rechner ähnlich fordern und der Vergleich das Modell misst, nicht die Größe.
# Alle vier tragen mindestens 128.000 Token Kontext, decken also jede Spalte ab.
MODELLE = [
    # Grundlinie: das derzeit eingesetzte Modell, Juli 2024
    ("mistral-nemo-instruct-2407", "Mistral NeMo 12B"),
    # Dieselbe Werkstatt, anderthalb Jahre später (Dezember 2025). Der
    # aussagekräftigste Vergleich - schlägt er die Grundlinie, ist der Umstieg
    # unmittelbar umsetzbar.
    ("ministral-3-14b-instruct-2512", "Ministral 3 14B"),
    # Juni 2026, gleiche Größenklasse, 256.000 Token Kontext
    ("gemma-4-12b-it", "Gemma 4 12B"),
    # Frühjahr 2026, etwas leichter - beantwortet zugleich, ob 12B nötig sind
    ("qwen3.5-9b", "Qwen3.5 9B"),
]
KONTEXTGROESSEN = [15000, 25000, 40000, 50000]

# Was für Anweisung, Auftrag und Antwort reserviert bleiben muss, in Token.
# Grundanweisung rund 800, Auftrag der Aufgabe rund 1100, Antwort bis 2600,
# dazu Sicherheitsabstand. Deutsche Fachtexte liegen bei etwa 3,3 Zeichen
# je Token.
RESERVE_TOKEN = 5500
ZEICHEN_JE_TOKEN = 3.3


# ---------------------------------------------------------------------------
# BEWERTUNG - die bekannte Wahrheit des Beispielfalls
# ---------------------------------------------------------------------------
EINZELPUNKTE = {1: 4, 2: 4, 3: 1, 4: 9, 5: 2, 6: 3}

# Woran sich erkennen lässt, dass ein Ansatzpunkt gefunden wurde. Geprüft wird
# im Textabschnitt zum jeweiligen Modul, nicht im ganzen Text - sonst zählte
# ein "erheblich" an ganz anderer Stelle mit.
ANSATZPUNKTE = {
    3: ("Verhaltensweisen", ("nächtlich", "unruh", "angst")),
    4: ("Selbstversorgung", ("täglich", "duschen", "zweite person", "ankleiden")),
    5: ("krankheits", ("insulin", "blutzucker", "dosierer", "verband", "medikament")),
}


@dataclass
class Ergebnis:
    modell: str
    kontext: int
    geladen: bool = False
    fehler: str = ""
    kontext_zeichen: int = 0
    belege: int = 0
    gegenbelege: int = 0
    sekunden_suche: float = 0.0
    sekunden_antwort: float = 0.0
    token_je_sekunde: float = 0.0
    antwort_zeichen: int = 0
    punkte: int = 0
    einzelbewertung: dict = field(default_factory=dict)
    antwort: str = ""

    @property
    def sekunden_gesamt(self) -> float:
        return self.sekunden_suche + self.sekunden_antwort


def bewerte(antwort: str, unterlagen: str) -> tuple[int, dict]:
    """Vergibt bis zu 15 Punkte gegen die bekannte Wahrheit des Falls."""
    import pflege_pdf

    klein = antwort.lower()
    einzeln: dict = {}
    punkte = 0

    # Die drei eingebauten Ansatzpunkte - je 2 Punkte
    for modul, (marke, belege) in ANSATZPUNKTE.items():
        stelle = klein.find(marke.lower())
        abschnitt = klein[stelle:stelle + 900] if stelle != -1 else ""
        erkannt = bool(abschnitt) and "erheblich" in abschnitt and any(
            b in abschnitt for b in belege
        )
        einzeln[f"Modul {modul} als erheblich erkannt"] = 2 if erkannt else 0
        punkte += 2 if erkannt else 0

    # Pflegegrad und Bescheiddatum - je 1 Punkt
    grad_ok = "pflegegrad 2" in klein and "pflegegrad 1" not in klein
    einzeln["Pflegegrad 2 richtig"] = 1 if grad_ok else 0
    punkte += 1 if grad_ok else 0

    datum_ok = "06.07.2026" in antwort or "6. juli 2026" in klein
    einzeln["Bescheiddatum richtig"] = 1 if datum_ok else 0
    punkte += 1 if datum_ok else 0

    # Einzelpunkte je Modul richtig abgelesen - insgesamt 3 Punkte
    richtig = sum(
        1 for wert in EINZELPUNKTE.values()
        if f"{wert} einzelpunkt" in klein or f"{wert} punkte" in klein
    )
    wert = round(3 * richtig / len(EINZELPUNKTE))
    einzeln[f"Einzelpunkte richtig ({richtig} von 6)"] = wert
    punkte += wert

    # Keine erfundenen Punktwerte - 2 Punkte
    erfunden = pflege_pdf.finde_unbelegte_punktzahlen(antwort, unterlagen)
    einzeln["keine erfundenen Punktzahlen"] = 0 if erfunden else 2
    punkte += 0 if erfunden else 2

    # Behauptet nicht, ein bewertetes Modul sei nicht bewertet worden - 1 Punkt
    falsch = any(w in klein for w in
                 ("nicht bewertet", "gar nicht bewertet", "fehlt eine bewertung"))
    einzeln["keine Falschbehauptung 'nicht bewertet'"] = 0 if falsch else 1
    punkte += 0 if falsch else 1

    # Unpersönlich formuliert - 1 Punkt
    persoenlich = pflege_pdf.find_personal_wording(antwort)
    einzeln["unpersönlich formuliert"] = 0 if persoenlich else 1
    punkte += 0 if persoenlich else 1

    return punkte, einzeln


# ---------------------------------------------------------------------------
# MODELLE LADEN
# ---------------------------------------------------------------------------
def lms(*argumente: str, timeout: int = 600) -> subprocess.CompletedProcess:
    return subprocess.run([str(LMS), *argumente], capture_output=True, text=True,
                          timeout=timeout, encoding="utf-8", errors="replace")


def passt_in_speicher(modell: str, kontext: int) -> tuple[bool, str]:
    """Fragt LM Studio vorab, ob die Kombination lädt - ohne sie zu laden."""
    ergebnis = lms("load", modell, "-c", str(kontext), "--estimate-only", "-y", timeout=180)
    ausgabe = (ergebnis.stdout + ergebnis.stderr).strip()
    passt = ergebnis.returncode == 0 and "not enough" not in ausgabe.lower()
    letzte = ausgabe.splitlines()[-1][:120] if ausgabe else ""
    return passt, letzte


def lade_modell(modell: str, kontext: int) -> tuple[bool, str]:
    import urllib.request

    lms("unload", "--all", timeout=120)
    time.sleep(2)
    ergebnis = lms("load", modell, "-c", str(kontext), "--gpu", "max", "-y", timeout=900)
    if ergebnis.returncode != 0:
        text = (ergebnis.stderr or ergebnis.stdout).strip()
        return False, text.splitlines()[-1][:160] if text else "Laden fehlgeschlagen"

    for _ in range(60):
        try:
            urllib.request.urlopen("http://127.0.0.1:1234/v1/models", timeout=5)
            return True, ""
        except Exception:
            time.sleep(2)
    return False, "Schnittstelle antwortet nach dem Laden nicht"


def grenzen_setzen(kontext: int) -> int:
    """Passt die Grenzen im Code an das Kontextfenster an.

    Ohne diesen Schritt wäre der Prompt in allen Spalten gleich groß und der
    ganze Vergleich träfe keine Aussage über die Wirkung des Fensters.
    """
    import pflege_rag

    budget_token = max(kontext - RESERVE_TOKEN, 2000)
    zeichen = int(budget_token * ZEICHEN_JE_TOKEN)
    pflege_rag.MAX_CONTEXT_CHARS = zeichen
    # Die Abschnittszahlen so hoch setzen, dass allein das Zeichenbudget
    # begrenzt - sonst wäre die Auswahl der Engpass statt des Fensters.
    pflege_rag.FINAL_USER_CHUNKS_BREIT = 400
    pflege_rag.FINAL_EXPERT_CHUNKS_BREIT = 400
    pflege_rag.BELEGE_JE_TEILFRAGE = max(2, budget_token // 1200)
    return zeichen


# ---------------------------------------------------------------------------
# EIN DURCHLAUF
# ---------------------------------------------------------------------------
def durchlauf(modell_key: str, anzeige: str, kontext: int) -> Ergebnis:
    import ingest
    import pflege_rag

    e = Ergebnis(modell=anzeige, kontext=kontext)
    ok, meldung = lade_modell(modell_key, kontext)
    if not ok:
        e.fehler = meldung
        return e
    e.geladen = True
    e.kontext_zeichen = grenzen_setzen(kontext)

    # 1. Unterlagen einlesen - jedes Mal neu, wie beim echten Hochladen
    schnell, mit_ocr = ingest.build_converters()
    dokumente = []
    for pfad in sorted(UNTERLAGEN.glob("*.pdf")):
        d = pflege_rag.extract_document_from_pdf(
            pfad.read_bytes(), pfad.name, schnell, ocr_converter=mit_ocr)
        if d is not None:
            dokumente.append(d)
    teile = pflege_rag.split_documents(dokumente)
    unterlagen_text = "\n".join(t.page_content for t in teile)

    # 2. Suchindex aufbauen
    einbett = pflege_rag.create_embeddings()
    speicher = pflege_rag.open_expert_database(einbett)
    fach = pflege_rag.HybridIndex(speicher, pflege_rag.load_all_expert_chunks(speicher))
    nutzer = pflege_rag.HybridIndex(
        pflege_rag.build_user_vector_store(teile, einbett), teile)
    reranker = pflege_rag.create_reranker()

    # 3. Differenzanalyse
    aktion = pflege_rag.QUICK_ACTION_BY_KEY["differenz"]
    t0 = time.time()
    ergebnis = pflege_rag.prepare_context(
        fach, nutzer, aktion.nutzertext, reranker=reranker,
        extra_queries=aktion.zusatzfragen)
    e.sekunden_suche = time.time() - t0
    e.belege = len(ergebnis.quellen)
    e.gegenbelege = sum(1 for q in ergebnis.quellen
                        if q.herkunft == "nutzer" and "Gutachten" not in q.quelle)

    if os.getenv("NACHEINANDER_LADEN", "0").strip() in ("1", "ja", "true"):
        pflege_rag.gib_reranker_frei(reranker)

    nachrichten = pflege_rag.build_messages(
        ergebnis.system_prompt,
        [{"role": "user", "content": aktion.render(
            pflege_rag.Antragsteller(versicherte_name="Hannelore Brandt"))}])

    t1 = time.time()
    stuecke: List[str] = []
    try:
        for teil in pflege_rag.stream_answer(pflege_rag.create_llm(), nachrichten):
            stuecke.append(teil)
    except Exception as fehler:
        e.fehler = f"{type(fehler).__name__}: {str(fehler)[:120]}"
    e.sekunden_antwort = time.time() - t1

    antwort = pflege_rag.render_citations(
        pflege_rag.strip_context_headers("".join(stuecke)), ergebnis.nummern)
    e.antwort = antwort
    e.antwort_zeichen = len(antwort)
    if e.sekunden_antwort > 0 and antwort:
        e.token_je_sekunde = len(antwort) / ZEICHEN_JE_TOKEN / e.sekunden_antwort

    # 4. Bewerten
    e.punkte, e.einzelbewertung = bewerte(antwort, unterlagen_text)
    return e


# ---------------------------------------------------------------------------
# AUSGABE
# ---------------------------------------------------------------------------
def schreibe_tabelle(ergebnisse: List[Ergebnis], ziel: str) -> None:
    """Schreibt die Kreuztabelle als Excel-Datei."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("openpyxl fehlt - bitte 'pip install openpyxl'. "
              "Die Ergebnisse liegen als vergleich_zwischenstand.json vor.")
        return

    mappe = Workbook()
    blatt = mappe.active
    blatt.title = "Vergleich"
    fett = Font(bold=True)
    kopffarbe = PatternFill("solid", fgColor="E8ECF0")

    kopf = blatt.cell(1, 1, "Modell")
    kopf.font, kopf.fill = fett, kopffarbe
    for spalte, kontext in enumerate(KONTEXTGROESSEN, start=2):
        z = blatt.cell(1, spalte, f"{kontext:,} Token".replace(",", "."))
        z.font, z.fill = fett, kopffarbe
        z.alignment = Alignment(horizontal="center")

    modelle = list(dict.fromkeys(e.modell for e in ergebnisse))
    for zeile, modell in enumerate(modelle, start=2):
        blatt.cell(zeile, 1, modell).font = fett
        for spalte, kontext in enumerate(KONTEXTGROESSEN, start=2):
            treffer = next((e for e in ergebnisse
                            if e.modell == modell and e.kontext == kontext), None)
            if treffer is None:
                text = "-"
            elif treffer.fehler and not treffer.antwort:
                text = f"nicht gelaufen\n{treffer.fehler[:60]}"
            else:
                text = (f"{treffer.punkte} von 15 Punkten\n"
                        f"{treffer.sekunden_gesamt:.0f} s gesamt\n"
                        f"{treffer.token_je_sekunde:.1f} Token/s\n"
                        f"{treffer.gegenbelege} Gegenbelege")
            z = blatt.cell(zeile, spalte, text)
            z.alignment = Alignment(wrap_text=True, vertical="top")

    blatt.column_dimensions["A"].width = 24
    for buchstabe in "BCDE":
        blatt.column_dimensions[buchstabe].width = 22
    for zeile in range(2, len(modelle) + 2):
        blatt.row_dimensions[zeile].height = 64

    detail = mappe.create_sheet("Einzelwerte")
    spalten = ["Modell", "Kontext", "Punkte", "Sekunden", "Token/s", "Belege",
               "Gegenbelege", "Kontextzeichen", "Antwortzeichen", "Fehler"]
    for i, name in enumerate(spalten, start=1):
        z = detail.cell(1, i, name)
        z.font, z.fill = fett, kopffarbe
    for zeile, e in enumerate(ergebnisse, start=2):
        werte = [e.modell, e.kontext, e.punkte, round(e.sekunden_gesamt),
                 round(e.token_je_sekunde, 1), e.belege, e.gegenbelege,
                 e.kontext_zeichen, e.antwort_zeichen, e.fehler]
        for i, wert in enumerate(werte, start=1):
            detail.cell(zeile, i, wert)
    for buchstabe, breite in zip("ABCDEFGHIJ", (24, 10, 8, 10, 10, 8, 12, 14, 14, 40)):
        detail.column_dimensions[buchstabe].width = breite

    mappe.save(ziel)
    print(f"\nTabelle geschrieben: {ziel}")


def main() -> int:
    zerteiler = argparse.ArgumentParser()
    zerteiler.add_argument("--nur-schaetzen", action="store_true",
                           help="nur prüfen, welche Kombinationen in den Speicher passen")
    zerteiler.add_argument("--ziel", default="vergleich_modelle.xlsx")
    argumente = zerteiler.parse_args()

    if not LMS.exists():
        print(f"lms nicht gefunden unter {LMS}")
        return 2
    if not UNTERLAGEN.is_dir():
        print(f"Unterlagen nicht gefunden unter {UNTERLAGEN}")
        return 2

    if argumente.nur_schaetzen:
        print(f"{'Modell':22s} {'Kontext':>8s}  passt?")
        for key, anzeige in MODELLE:
            for kontext in KONTEXTGROESSEN:
                passt, meldung = passt_in_speicher(key, kontext)
                print(f"{anzeige:22s} {kontext:8d}  {'ja  ' if passt else 'NEIN'}  {meldung}")
        return 0

    ergebnisse: List[Ergebnis] = []
    gesamt = len(MODELLE) * len(KONTEXTGROESSEN)
    nummer = 0
    for key, anzeige in MODELLE:
        for kontext in KONTEXTGROESSEN:
            nummer += 1
            print(f"\n[{nummer}/{gesamt}] {anzeige} bei {kontext} Token", flush=True)
            start = time.time()
            try:
                e = durchlauf(key, anzeige, kontext)
            except Exception as fehler:
                e = Ergebnis(modell=anzeige, kontext=kontext,
                             fehler=f"{type(fehler).__name__}: {str(fehler)[:140]}")
            ergebnisse.append(e)
            if e.fehler and not e.antwort:
                print(f"    nicht gelaufen: {e.fehler}")
            else:
                print(f"    {e.punkte} von 15 Punkten, {e.sekunden_gesamt:.0f}s, "
                      f"{e.token_je_sekunde:.1f} Token/s, {e.gegenbelege} Gegenbelege")
            # Zwischenstand sichern - der ganze Lauf dauert Stunden
            Path("vergleich_zwischenstand.json").write_text(
                json.dumps([asdict(x) for x in ergebnisse], ensure_ascii=False, indent=1),
                encoding="utf-8")
            print(f"    (Teilzeit {time.time() - start:.0f}s)", flush=True)

    schreibe_tabelle(ergebnisse, argumente.ziel)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
